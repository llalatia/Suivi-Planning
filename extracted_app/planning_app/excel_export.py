# -*- coding: utf-8 -*-
"""
excel_export.py — Régénère un fichier Excel "Macro-Planning" à partir de la
base SQLite, avec la même mise en page que le fichier d'origine :
  - ligne 1 : titre
  - ligne 2 : code de la semaine courante
  - ligne 3 : dates de chaque semaine (fusionnées sur 3 colonnes)
  - ligne 4 : codes semaine (fusionnés sur 3 colonnes)
  - ligne 5 : en-têtes de colonnes
  - lignes suivantes : une ligne par commande
  - légende couleurs en bas de page

Les valeurs sont écrites en dur (pas de formules Excel) car elles sont
calculées par l'application elle-même (logic.py) — l'export est un rapport,
pas un modèle à modifier à la main.
"""

import datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

import db
import logic

FONT_NAME = "Calibri"
COLOR_HEADER_FILL = "FFD9D9D9"
COLOR_SEMAINE_COURANTE = "FFFFFF00"   # jaune
COLOR_LANCEMENT = "FFFFC000"          # orange
COLOR_LIMITE_CLIENT = "FFFF0000"      # rouge
COLOR_SOUS_PLANIFIE = "FFFFC7CE"      # rose clair (alerte, texte)
COLOR_TERMINE = "FFC6EFCE"            # vert clair

THIN = Side(style="thin", color="FFB7B7B7")
BORDER_ALL = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def _fmt_num(v):
    if v is None:
        return None
    if isinstance(v, float) and v.is_integer():
        return int(v)
    return v


def export_xlsx(path, semaine_courante_code=None, titre="Planning BRODERIE TRADITIONNELLE"):
    semaines = db.list_semaines()
    commandes = db.list_commandes()

    if semaine_courante_code is None:
        semaine_courante_code = db.get_config("semaine_courante")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"MAJ {semaine_courante_code or ''}".strip()

    bold = Font(name=FONT_NAME, bold=True)
    normal = Font(name=FONT_NAME)

    # --- Ligne 1-2 : titre / semaine courante ---
    ws["B1"] = titre
    ws["B1"].font = Font(name=FONT_NAME, bold=True, size=12)
    ws.merge_cells("B1:E1")

    ws["B2"] = semaine_courante_code or ""
    ws["B2"].font = bold
    if semaine_courante_code:
        ws["B2"].fill = PatternFill("solid", fgColor=COLOR_SEMAINE_COURANTE)

    # --- Colonnes fixes ---
    fixed_headers = ["CLIENT", "MODELE", "QTE EN\nCOMMANDE", "UNITE",
                      "HEURE / PCS", "QTE FINIE", "QTE RESTE\n A FAIRE",
                      "PIECES PROGRAMMES DANS LE PLANNING"]
    for i, h in enumerate(fixed_headers):
        col = i + 1
        c = ws.cell(row=5, column=col, value=h)
        c.font = bold
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        c.fill = PatternFill("solid", fgColor=COLOR_HEADER_FILL)
        c.border = BORDER_ALL

    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 40
    ws.column_dimensions["C"].width = 14
    ws.column_dimensions["D"].width = 9
    ws.column_dimensions["E"].width = 14
    ws.column_dimensions["F"].width = 11
    ws.column_dimensions["G"].width = 14
    ws.column_dimensions["H"].width = 16

    # --- Colonnes par semaine (3 colonnes chacune : brodeuses / heures / pieces) ---
    start_col = 9  # colonne I
    col_for_semaine = {}
    for idx, s in enumerate(semaines):
        c0 = start_col + idx * 3
        col_for_semaine[s["id"]] = c0

        # ligne 3 : date (fusionnée sur 3 colonnes)
        try:
            d = datetime.datetime.strptime(s["date_debut"], "%Y-%m-%d")
            date_cell = ws.cell(row=3, column=c0, value=d)
            date_cell.number_format = "dd/mm/yyyy"
        except ValueError:
            ws.cell(row=3, column=c0, value=s["date_debut"])
        ws.merge_cells(start_row=3, start_column=c0, end_row=3, end_column=c0 + 2)
        ws.cell(row=3, column=c0).font = bold
        ws.cell(row=3, column=c0).alignment = Alignment(horizontal="center")

        # ligne 4 : code semaine (fusionné)
        wk_cell = ws.cell(row=4, column=c0, value=s["code"])
        ws.merge_cells(start_row=4, start_column=c0, end_row=4, end_column=c0 + 2)
        wk_cell.font = bold
        wk_cell.alignment = Alignment(horizontal="center")

        is_courante = (s["code"] == semaine_courante_code)
        if is_courante:
            for cc in (c0, c0 + 1, c0 + 2):
                ws.cell(row=3, column=cc).fill = PatternFill("solid", fgColor=COLOR_SEMAINE_COURANTE)
                ws.cell(row=4, column=cc).fill = PatternFill("solid", fgColor=COLOR_SEMAINE_COURANTE)
                ws.cell(row=5, column=cc).fill = PatternFill("solid", fgColor=COLOR_SEMAINE_COURANTE)

        # ligne 5 : en-têtes
        labels = ["Nb brodeuses", "Nb d'heures", "Nb pièces "]
        for j, lbl in enumerate(labels):
            hc = ws.cell(row=5, column=c0 + j, value=lbl)
            hc.font = bold
            hc.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
            hc.border = BORDER_ALL
            if not is_courante:
                hc.fill = PatternFill("solid", fgColor=COLOR_HEADER_FILL)
        ws.column_dimensions[get_column_letter(c0)].width = 11.5
        ws.column_dimensions[get_column_letter(c0 + 1)].width = 11.5
        ws.column_dimensions[get_column_letter(c0 + 2)].width = 11.5

    ws.freeze_panes = "C6"

    # --- Lignes de données ---
    row = 6
    for cmd in commandes:
        c = dict(cmd)
        ws.cell(row=row, column=1, value=c["client"])
        ws.cell(row=row, column=2, value=c["modele"])
        ws.cell(row=row, column=3, value=_fmt_num(c["qte_commande"]))
        ws.cell(row=row, column=4, value=c["unite"])
        ws.cell(row=row, column=5, value="TBC" if c["heure_pcs_tbc"] else _fmt_num(c["heure_pcs"]))
        ws.cell(row=row, column=6, value=_fmt_num(c["qte_finie"]))

        reste = logic.qte_reste_a_faire(c["qte_commande"], c["qte_finie"])
        ws.cell(row=row, column=7, value=_fmt_num(reste))

        planning = db.planning_for_commande(c["id"])
        total_programme = 0
        for s in semaines:
            c0 = col_for_semaine[s["id"]]
            p = planning.get(s["id"])
            pieces = p["nb_pieces"] if p else 0
            statut = p["statut"] if p else "normal"
            total_programme += pieces or 0

            heures = logic.nb_heures(pieces, c["heure_pcs"]) if not c["heure_pcs_tbc"] else None
            brodeuses = logic.nb_brodeuses(heures)

            cell_b = ws.cell(row=row, column=c0, value=_fmt_num(brodeuses))
            cell_h = ws.cell(row=row, column=c0 + 1, value=_fmt_num(heures))
            cell_p = ws.cell(row=row, column=c0 + 2, value=_fmt_num(pieces) if pieces else None)

            for cell in (cell_b, cell_h, cell_p):
                cell.border = BORDER_ALL
                if statut == "lancement":
                    cell.fill = PatternFill("solid", fgColor=COLOR_LANCEMENT)
                elif statut == "limite_client":
                    cell.fill = PatternFill("solid", fgColor=COLOR_LIMITE_CLIENT)
                elif s["code"] == semaine_courante_code:
                    cell.fill = PatternFill("solid", fgColor="FFFFFAE0")

        ws.cell(row=row, column=8, value=_fmt_num(total_programme))

        # Mise en forme conditionnelle simple (statut global de la commande)
        st = logic.statut_commande(c["qte_commande"], c["qte_finie"], total_programme)
        row_fill = None
        if st == "termine":
            row_fill = COLOR_TERMINE
        elif st == "sous_planifie":
            row_fill = COLOR_SOUS_PLANIFIE
        if row_fill:
            for col in range(1, 9):
                existing = ws.cell(row=row, column=col)
                if existing.fill.fgColor.rgb in (None, "00000000"):
                    existing.fill = PatternFill("solid", fgColor=row_fill)

        for col in range(1, 9):
            ws.cell(row=row, column=col).font = normal
            ws.cell(row=row, column=col).border = BORDER_ALL

        row += 1

    # --- Ligne de totaux ---
    total_row = row + 1
    ws.cell(row=total_row, column=8, value="TOTAUX SEMAINE").font = bold
    for s in semaines:
        c0 = col_for_semaine[s["id"]]
        total_pieces = 0
        total_heures = 0
        with db.get_conn() as conn:
            recs = conn.execute(
                "SELECT p.nb_pieces, c.heure_pcs, c.heure_pcs_tbc FROM planning p "
                "JOIN commandes c ON c.id = p.commande_id WHERE p.semaine_id=?",
                (s["id"],),
            ).fetchall()
        for r_ in recs:
            pieces = r_["nb_pieces"] or 0
            total_pieces += pieces
            if not r_["heure_pcs_tbc"] and r_["heure_pcs"]:
                total_heures += pieces * r_["heure_pcs"]
        total_brodeuses = logic.nb_brodeuses(total_heures)
        ws.cell(row=total_row, column=c0, value=_fmt_num(total_brodeuses)).font = bold
        ws.cell(row=total_row, column=c0 + 1, value=_fmt_num(round(total_heures, 2))).font = bold
        ws.cell(row=total_row, column=c0 + 2, value=_fmt_num(total_pieces)).font = bold

    # --- Légende ---
    legend_row = total_row + 3
    ws.cell(row=legend_row, column=2, value="LEGENDE").font = bold
    legend_items = [
        ("Semaine en cours", COLOR_SEMAINE_COURANTE),
        ("Semaine de lancement", COLOR_LANCEMENT),
        ("Limite livraison client", COLOR_LIMITE_CLIENT),
        ("Commande sous-planifiée (retard)", COLOR_SOUS_PLANIFIE),
        ("Commande terminée", COLOR_TERMINE),
    ]
    for i, (label, color) in enumerate(legend_items):
        r_ = legend_row + 1 + i
        ws.cell(row=r_, column=2, value=label).font = normal
        ws.cell(row=r_, column=1).fill = PatternFill("solid", fgColor=color)
        ws.cell(row=r_, column=1).border = BORDER_ALL

    wb.save(path)
    return path


if __name__ == "__main__":
    import sys
    out = sys.argv[1] if len(sys.argv) > 1 else "export_test.xlsx"
    sem = sys.argv[2] if len(sys.argv) > 2 else None
    export_xlsx(out, semaine_courante_code=sem)
    print("Export écrit dans", out)
